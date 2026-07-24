# -*- coding: utf-8 -*-
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUT = Path("outputs/aigc_detection_provenance_2026_monthly_report.pptx")

WIDE_W = 13.333
WIDE_H = 7.5

COLORS = {
    "ink": "111827",
    "ink2": "1F2937",
    "muted": "64748B",
    "light": "F8FAFC",
    "line": "CBD5E1",
    "teal": "0F766E",
    "teal2": "14B8A6",
    "amber": "F59E0B",
    "orange": "EA580C",
    "red": "DC2626",
    "green": "16A34A",
    "blue": "2563EB",
    "violet": "7C3AED",
    "white": "FFFFFF",
    "panel": "EEF6F5",
}


def rgb(hex_color):
    hex_color = hex_color.replace("#", "")
    return RGBColor(int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16))


def set_fill(shape, color, transparency=0):
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)
    if transparency:
        shape.fill.transparency = transparency


def set_line(shape, color=None, width=1, transparency=0):
    if color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = rgb(color)
        shape.line.width = Pt(width)
        if transparency:
            shape.line.transparency = transparency


def add_bg(slide, color="F8FAFC"):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(WIDE_W), Inches(WIDE_H))
    set_fill(bg, color)
    set_line(bg, None)
    bg.z_order = 0


def add_text(slide, text, x, y, w, h, size=18, color="111827", bold=False, align="left",
             valign="top", font="Microsoft YaHei", margin=0.04):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[valign]
    p = tf.paragraphs[0]
    p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return box


def add_multiline(slide, lines, x, y, w, h, size=15, color="111827", bullet=False,
                  font="Microsoft YaHei", gap=2):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.06)
    tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.03)
    tf.margin_bottom = Inches(0.03)
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = font
        p.font.size = Pt(size)
        p.font.color.rgb = rgb(color)
        p.space_after = Pt(gap)
        if bullet:
            p.level = 0
            p._p.get_or_add_pPr().set("marL", "220000")
            p._p.get_or_add_pPr().set("indent", "-120000")
    return box


def add_title(slide, title, subtitle=None, section=None, slide_no=None, dark=False):
    color = "FFFFFF" if dark else "111827"
    muted = "CBD5E1" if dark else "64748B"
    if section:
        add_text(slide, section, 0.58, 0.28, 2.8, 0.26, 9, muted, True)
    add_text(slide, title, 0.58, 0.54, 11.2, 0.52, 28, color, True)
    if subtitle:
        add_text(slide, subtitle, 0.60, 1.05, 11.0, 0.32, 12, muted)
    if slide_no is not None:
        add_text(slide, f"{slide_no:02d}", 12.35, 7.02, 0.45, 0.2, 8, muted, True, "right")


def add_card(slide, x, y, w, h, title, body=None, color="FFFFFF", accent="0F766E",
             title_size=16, body_size=12, shadow=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    set_fill(shape, color)
    set_line(shape, "D7DEE8", 0.8)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.07), Inches(h))
    set_fill(bar, accent)
    set_line(bar, None)
    if body is None:
        add_text(slide, title, x + 0.18, y + 0.09, w - 0.32, h - 0.18, title_size, "111827", True, valign="middle")
        return shape
    title_h = 0.56 if "\n" in title else 0.28
    add_text(slide, title, x + 0.18, y + 0.12, w - 0.32, title_h, title_size, "111827", True)
    if body:
        body_y = y + 0.30 + title_h
        body_h = h - title_h - 0.42
        if isinstance(body, list):
            add_multiline(slide, body, x + 0.18, body_y, w - 0.32, body_h, body_size, "334155")
        else:
            add_text(slide, body, x + 0.18, body_y, w - 0.32, body_h, body_size, "334155")
    return shape


def add_badge(slide, text, x, y, w, color="0F766E", text_color="FFFFFF"):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.32))
    set_fill(shape, color)
    set_line(shape, None)
    add_text(slide, text, x + 0.03, y + 0.055, w - 0.06, 0.18, 8.5, text_color, True, "center")
    return shape


def add_circle_label(slide, text, x, y, d=0.42, color="0F766E"):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    set_fill(shape, color)
    set_line(shape, None)
    add_text(slide, text, x, y + 0.08, d, d - 0.14, 11, "FFFFFF", True, "center")
    return shape


def add_arrow(slide, x1, y1, x2, y2, color="94A3B8", width=1.5):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    line.line.color.rgb = rgb(color)
    line.line.width = Pt(width)
    line.line.end_arrowhead = True
    return line


def add_source(slide, text):
    add_text(slide, text, 0.58, 7.06, 10.8, 0.18, 7.2, "64748B")


def slide_cover(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, "111827")
    # Subtle geometry
    for i, (x, y, w, h, c) in enumerate([
        (8.15, 0.05, 4.85, 1.75, "0F766E"),
        (8.95, 1.85, 4.10, 1.90, "2563EB"),
        (8.35, 4.10, 4.70, 1.85, "F59E0B"),
    ]):
        shp = slide.shapes.add_shape(MSO_SHAPE.PARALLELOGRAM, Inches(x), Inches(y), Inches(w), Inches(h))
        set_fill(shp, c, 26 + i * 8)
        set_line(shp, None)
    add_badge(slide, "月度阶段汇报 2026.07", 0.70, 0.68, 1.75, "0F766E")
    add_text(slide, "AIGC 检测与溯源系统", 0.70, 1.42, 7.8, 0.7, 36, "FFFFFF", True)
    add_text(slide, "从可运行框架到可部署、可扩展、可验证", 0.73, 2.16, 7.8, 0.36, 19, "D1D5DB", True)
    add_text(slide, "服务器部署  |  2026 前沿调研  |  多模态样本扩展", 0.75, 2.74, 7.2, 0.32, 15, "CBD5E1")

    # Right flow visual
    labels = [("输入内容", "文本/图像/音频/视频"), ("多证据检测", "被动检测 + 水印 + C2PA"), ("溯源报告", "Top-K 来源 + 证据链")]
    y = 1.25
    for idx, (t, b) in enumerate(labels):
        add_card(slide, 8.55, y + idx * 1.45, 3.25, 0.92, t, b, color="F8FAFC", accent=["14B8A6", "F59E0B", "2563EB"][idx], title_size=15, body_size=10)
        if idx < 2:
            add_arrow(slide, 10.18, y + idx * 1.45 + 0.93, 10.18, y + (idx + 1) * 1.45 - 0.05, "D1D5DB", 1.5)
    add_text(slide, "汇报重点：2026 年真实场景泛化与双层溯源标识", 0.75, 6.72, 7.6, 0.22, 11, "9CA3AF")
    return slide


def slide_overview(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_title(slide, "本月工作总览", "三条线同时推进：部署交付、样本扩展、2026 前沿调研", "01 / OVERVIEW", 2)
    lanes = [
        ("工程交付", "Linux 服务器部署文档", ["系统依赖", "Python 环境", "模型缓存", "API Key", "C2PA / 水印工具"], "0F766E"),
        ("样本扩展", "为检测与归因准备新模型样本", ["文本 JSONL", "图片分目录", "音视频片段", "来源模型标签"], "F59E0B"),
        ("前沿调研", "聚焦 2026 最新变化", ["in-the-wild", "跨模型泛化", "多语言任务", "C2PA 安全争议"], "2563EB"),
    ]
    for i, (name, title, items, c) in enumerate(lanes):
        x = 0.75 + i * 4.05
        add_circle_label(slide, str(i + 1), x, 1.73, 0.46, c)
        add_card(slide, x, 2.28, 3.55, 3.6, name, [title, "", *items], color="FFFFFF", accent=c, title_size=17, body_size=12)
        tick = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, Inches(x + 0.23), Inches(2.18), Inches(x + 0.23), Inches(2.65)
        )
        tick.line.color.rgb = rgb(c)
        tick.line.width = Pt(1.8)
    add_text(slide, "这次汇报不再重复“检测层有哪些模型”，而是回答：2026 年新问题如何影响系统下一步设计。", 0.78, 6.46, 11.7, 0.34, 14, "334155", True)
    return slide


def slide_before_after(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, "F8FAFC")
    add_title(slide, "从上次架构设计到本次落地路线", "上次解决“系统怎么拆”，这次解决“如何部署、如何评测、如何跟上 2026 新场景”。", "02 / CONTEXT", 3)
    add_card(slide, 0.78, 1.70, 5.25, 4.25, "上次汇报：模块化架构", [
        "内容接入与预处理",
        "多模态检测引擎",
        "C2PA / 水印 / 指纹",
        "Top-K 来源归因",
        "SQLite + API + Dashboard",
    ], "FFFFFF", "64748B", 18, 14)
    add_arrow(slide, 6.15, 3.65, 7.10, 3.65, "0F766E", 2.6)
    add_card(slide, 7.28, 1.70, 5.25, 4.25, "本次汇报：可部署和可验证", [
        "Linux 部署文档已形成",
        "样本库扩展到新生成模型",
        "引入 2026 benchmark 思路",
        "强调真实传播场景鲁棒性",
        "C2PA + 水印 + 检测融合",
    ], "FFFFFF", "0F766E", 18, 14)
    add_text(slide, "核心变化", 5.68, 2.62, 1.70, 0.26, 13, "0F766E", True, "center")
    add_badge(slide, "从“能跑通”到“能解释、可部署、可更新”", 4.55, 6.34, 4.25, "111827")
    return slide


def slide_deployment(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_title(slide, "Linux 部署文档交付", "部署文档按“先跑通，再增强”组织，覆盖从环境到模型、工具、服务启动的完整链路。", "03 / DELIVERY", 4)
    steps = [
        ("01", "系统依赖", "git / ffmpeg / libsndfile / libGL"),
        ("02", "Python 环境", "venv + requirements + datasets/joblib"),
        ("03", "缓存目录", "HF_HOME / HF_HUB_CACHE / DATASETS_CACHE"),
        ("04", "配置安全", "API Key 改为环境变量名"),
        ("05", "增强工具", "c2patool / AudioSeal / VideoSeal"),
        ("06", "服务启动", "FastAPI 8000 + Streamlit 8501"),
    ]
    for i, (num, t, b) in enumerate(steps):
        x = 0.72 + (i % 3) * 4.05
        y = 1.55 + (i // 3) * 2.25
        add_card(slide, x, y, 3.45, 1.45, t, b, "FFFFFF", ["0F766E", "14B8A6", "F59E0B", "2563EB", "7C3AED", "DC2626"][i], 15, 11)
        add_badge(slide, num, x + 2.67, y + 0.12, 0.45, ["0F766E", "14B8A6", "F59E0B", "2563EB", "7C3AED", "DC2626"][i])
    add_text(slide, "部署结论：基础 API 和 Dashboard 可先上线，C2PA / 水印 / 来源归因按增强项逐步补齐。", 0.78, 6.55, 11.8, 0.28, 14, "334155", True)
    add_source(slide, "来源：docs/LINUX_DEPLOYMENT.md")
    return slide


def slide_status(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, "F8FAFC")
    add_title(slide, "当前系统落地状态", "系统已经是可运行框架，但真实检测 provider、官方 C2PA、水印模型仍需要逐步接入。", "04 / STATUS", 5)
    headers = ["模块", "当前状态", "下一步"]
    rows = [
        ("FastAPI / Streamlit", "真实可运行", "保持接口稳定"),
        ("Ingestion / 指纹 / SQLite", "真实可运行", "增加 MIME sniffing"),
        ("Report", "模板报告可用", "加入证据限制说明"),
        ("Detection Provider", "demo_api 可演示", "接 Hive 或 Sightengine"),
        ("C2PA", "placeholder parser", "接 c2patool wrapper"),
        ("Watermark", "adapter 已预留", "接 AudioSeal / VideoSeal"),
    ]
    x0, y0 = 0.75, 1.55
    colw = [3.05, 3.75, 4.6]
    for j, h in enumerate(headers):
        shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x0 + sum(colw[:j])), Inches(y0), Inches(colw[j]), Inches(0.52))
        set_fill(shp, "111827")
        set_line(shp, "111827")
        add_text(slide, h, x0 + sum(colw[:j]) + 0.08, y0 + 0.15, colw[j] - 0.16, 0.18, 11, "FFFFFF", True)
    for i, row in enumerate(rows):
        y = y0 + 0.52 + i * 0.67
        for j, cell in enumerate(row):
            shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x0 + sum(colw[:j])), Inches(y), Inches(colw[j]), Inches(0.67))
            set_fill(shp, "FFFFFF" if i % 2 == 0 else "F1F5F9")
            set_line(shp, "E2E8F0", 0.6)
            color = "0F766E" if ("真实" in cell or "可用" in cell) else "334155"
            add_text(slide, cell, x0 + sum(colw[:j]) + 0.08, y + 0.18, colw[j] - 0.16, 0.24, 10.5, color, j == 1 and ("真实" in cell or "可用" in cell))
    add_badge(slide, "汇报口径：真实说清楚，预留也说清楚", 4.65, 6.33, 3.55, "0F766E")
    add_source(slide, "来源：docs/IMPLEMENTATION_STATUS.md")
    return slide


def slide_why_2026(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, "111827")
    add_title(slide, "为什么本月要重点调研 2026", "生成模型更新太快，旧检测器在真实场景下的泛化风险正在放大。", "05 / WHY NOW", 6, True)
    stats = [
        ("291", "unique generators", "开箱图像 detector benchmark"),
        ("42", "image generators", "NTIRE 2026 in-the-wild challenge"),
        ("36", "real transforms", "压缩、裁剪、模糊、缩放等"),
        ("10-40%", "accuracy drop", "真实 Wikipedia 任务型文本检测"),
    ]
    for i, (num, label, note) in enumerate(stats):
        x = 0.75 + i * 3.15
        shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.82), Inches(2.55), Inches(2.55))
        set_fill(shp, "FFFFFF", 8)
        set_line(shp, "334155", 0.8)
        add_text(slide, num, x + 0.12, 2.17, 2.28, 0.55, 34, ["14B8A6", "F59E0B", "60A5FA", "F87171"][i], True, "center")
        add_text(slide, label, x + 0.12, 2.88, 2.28, 0.28, 12, "E5E7EB", True, "center")
        add_text(slide, note, x + 0.18, 3.35, 2.16, 0.44, 9.5, "CBD5E1", align="center")
    add_text(slide, "2026 年新问题：不再只是“是不是 AI”，而是“传播后还能不能识别、能不能解释、能不能追踪来源”。", 1.25, 5.28, 10.9, 0.5, 19, "FFFFFF", True, "center")
    add_source(slide, "来源：NTIRE 2026、out-of-box image detector benchmark、TSM-Bench")
    return slide


def slide_image_2026(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_title(slide, "2026 图像检测新结论：真实变换成为主战场", "NTIRE 2026 把检测目标从干净数据集推进到真实传播场景。", "06 / IMAGE", 7)
    add_card(slide, 0.75, 1.52, 4.3, 4.7, "NTIRE 2026 Challenge", [
        "108,750 真实图像",
        "185,750 AI 生成图像",
        "42 个开放/闭源生成器",
        "36 种真实图像变换",
        "以 ROC AUC 评估鲁棒性",
    ], "FFFFFF", "0F766E", 18, 13)
    # Transformation wheel
    center_x, center_y = 8.6, 3.65
    labels = [("压缩", 7.2, 1.7, "F59E0B"), ("裁剪", 9.7, 1.72, "2563EB"), ("模糊", 10.35, 3.7, "7C3AED"), ("缩放", 9.1, 5.25, "DC2626"), ("截图", 6.95, 4.95, "16A34A"), ("滤镜", 6.45, 3.0, "14B8A6")]
    core = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(center_x - 0.9), Inches(center_y - 0.9), Inches(1.8), Inches(1.8))
    set_fill(core, "111827")
    set_line(core, None)
    add_text(slide, "in-the-wild\nrobustness", center_x - 0.72, center_y - 0.22, 1.44, 0.55, 12, "FFFFFF", True, "center")
    for txt, x, y, c in labels:
        add_arrow(slide, center_x, center_y, x + 0.33, y + 0.18, "94A3B8", 1)
        add_badge(slide, txt, x, y, 0.68, c)
    add_text(slide, "PPT 表达：图片检测不能只展示上传原图的准确率，要展示压缩、裁剪、截图后的稳定性。", 6.05, 6.18, 6.4, 0.36, 13, "334155", True)
    add_source(slide, "来源：https://arxiv.org/abs/2604.11487")
    return slide


def slide_out_of_box(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, "F8FAFC")
    add_title(slide, "2026 图像检测新结论：开箱 detector 没有通用赢家", "大规模 benchmark 显示，不同数据集之间 detector 排名极不稳定。", "07 / BENCHMARK", 8)
    bars = [
        ("最佳平均准确率", 75.0, "0F766E"),
        ("最差平均准确率", 37.5, "DC2626"),
        ("商业新模型场景", 30.0, "F59E0B"),
    ]
    x0, y0 = 1.0, 1.85
    for i, (label, val, c) in enumerate(bars):
        y = y0 + i * 1.05
        add_text(slide, label, x0, y + 0.10, 2.2, 0.24, 12, "334155", True)
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x0 + 2.25), Inches(y), Inches(5.6), Inches(0.42))
        set_fill(bg, "E2E8F0")
        set_line(bg, None)
        fg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x0 + 2.25), Inches(y), Inches(5.6 * val / 100), Inches(0.42))
        set_fill(fg, c)
        set_line(fg, None)
        add_text(slide, f"{val:.1f}%", x0 + 8.0, y + 0.05, 0.8, 0.24, 12, c, True)
    add_card(slide, 9.6, 1.55, 2.65, 3.88, "部署启发", [
        "不要绑定一个图像 detector",
        "按目标场景选择模型",
        "保留 provider ensemble",
        "持续补充新生成器样本",
    ], "FFFFFF", "2563EB", 16, 12)
    add_text(slide, "可放 PPT 的一句话：开源 detector 的论文分数不能直接等价于生产可用，必须按真实威胁模型重新评测。", 1.0, 6.26, 11.2, 0.34, 14, "111827", True)
    add_source(slide, "来源：https://arxiv.org/abs/2602.07814")
    return slide


def slide_text_2026(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_title(slide, "2026 文本检测新结论：真实任务比普通生成更难", "多语言、真实编辑任务、润色改写，会让传统文本检测器明显失效。", "08 / TEXT", 9)
    items = [
        ("DetectRL-X", "8 种语言、6 个高风险商业场景、4 个商业 LLM，并加入润色/扩写/缩写/改写攻击。", "0F766E"),
        ("TSM-Bench", "Wikipedia 真实编辑任务下，检测准确率比传统 benchmark 下降 10-40%。", "F59E0B"),
        ("Detecting the Machine", "域内表现可以很好，但跨域、跨 LLM、对抗条件下泛化下降；LLM-as-detector 也不稳定。", "2563EB"),
        ("LLM-specific fine-tuning", "按单个 LLM 或模型族训练 detector 能提升性能，但带来持续维护成本。", "7C3AED"),
    ]
    for i, (t, b, c) in enumerate(items):
        x = 0.75 + (i % 2) * 6.1
        y = 1.55 + (i // 2) * 2.05
        add_card(slide, x, y, 5.45, 1.5, t, b, "FFFFFF", c, 16, 11)
    # sample schema strip
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.05), Inches(6.05), Inches(11.1), Inches(0.72))
    set_fill(shp, "111827")
    set_line(shp, None)
    add_text(slide, "样本 schema 要新增：language、source_model、operation、domain、attack_type", 1.30, 6.27, 10.55, 0.20, 13, "FFFFFF", True, "center")
    add_source(slide, "来源：arXiv 2605.15518 / 2605.31113 / 2603.17522 / 2601.20006")
    return slide


def slide_social_vlm(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, "F8FAFC")
    add_title(slide, "2026 社交媒体场景：多模态 VLM + 可解释检测", "新工作强调在动态社交平台中持续采集数据，并用紧凑 VLM 同时做检测和解释。", "09 / SOCIAL", 10)
    stages = [
        ("社交平台内容", "图像 / 视频 / 文本 / 音频"),
        ("持续数据策展", "新模型、新模板、新传播方式"),
        ("紧凑 VLM", "多模态检测 + 解释生成"),
        ("平台反馈", "推荐、审核、事实核查"),
    ]
    for i, (t, b) in enumerate(stages):
        x = 0.78 + i * 3.12
        add_card(slide, x, 2.15, 2.45, 1.65, t, b, "FFFFFF", ["0F766E", "F59E0B", "2563EB", "DC2626"][i], 15, 11)
        if i < len(stages) - 1:
            add_arrow(slide, x + 2.48, 2.98, x + 2.95, 2.98, "64748B", 1.6)
    add_card(slide, 1.2, 5.0, 10.9, 0.95, "对本项目的直接启发", "Dashboard 不应只显示 AI 概率，还要展示可解释证据：可疑帧、可疑片段、命中水印、C2PA 状态、Top-K 来源。", "FFFFFF", "111827", 16, 13)
    add_source(slide, "来源：https://arxiv.org/abs/2606.11200")
    return slide


def slide_samples(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_title(slide, "样本扩展方案：从“AI/人类”到“来源模型 + 任务 + 传播变换”", "样本库要同时服务二分类检测、Top-K 来源归因和鲁棒性评测。", "10 / DATA", 11)
    rows = [
        ("文本", "ChatGPT / DeepSeek / Kimi / Claude / Gemini / Qwen", "语言、领域、任务、改写攻击"),
        ("图片", "GPT Image / Midjourney v7 / Flux / Firefly v4 / SDXL", "来源模型、压缩、裁剪、截图"),
        ("音频", "OpenAI TTS / ElevenLabs / Suno / Udio / CosyVoice", "片段级标签、水印、语音/音乐"),
        ("视频", "Veo / Sora / Runway / Kling / Pika", "抽帧证据、音画一致性、片段定位"),
    ]
    for i, (mod, models, fields) in enumerate(rows):
        y = 1.55 + i * 1.15
        add_badge(slide, mod, 0.85, y + 0.10, 0.85, ["0F766E", "F59E0B", "2563EB", "7C3AED"][i])
        add_card(slide, 1.88, y, 5.25, 0.78, models, None, "FFFFFF", ["0F766E", "F59E0B", "2563EB", "7C3AED"][i], 11, 10)
        add_card(slide, 7.35, y, 4.95, 0.78, fields, None, "FFFFFF", ["0F766E", "F59E0B", "2563EB", "7C3AED"][i], 11, 10)
    add_text(slide, "关键点：样本生成不要只用“请写一篇文章”，要包含摘要、润色、扩写、缩写、改写和平台传播后的版本。", 0.88, 6.48, 11.4, 0.25, 13, "334155", True)
    return slide


def slide_provenance_security(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, "111827")
    add_title(slide, "2026 溯源争议：C2PA 是强证据，但不能单独依赖", "独立安全分析提醒：内容凭证是重要方向，但在高风险场景仍需要多证据交叉验证。", "11 / PROVENANCE", 12, True)
    add_card(slide, 0.85, 1.70, 5.55, 3.55, "C2PA 的价值", [
        "记录生成/编辑/签名信息",
        "可用 c2patool 读取与验证",
        "适合形成可解释 provenance 证据",
        "有则价值高，无则不能反推不是 AI",
    ], "F8FAFC", "0F766E", 18, 13)
    add_card(slide, 6.95, 1.70, 5.55, 3.55, "2026 安全提醒", [
        "规范仍存在安全目标不足",
        "社交平台可能剥离 metadata",
        "截图/再编码会破坏凭证链",
        "不能作为最终判定的唯一依据",
    ], "F8FAFC", "DC2626", 18, 13)
    add_text(slide, "系统策略：C2PA reader + watermark adapter + passive detector + Top-K attribution + fingerprint registry", 1.2, 6.10, 10.9, 0.32, 15, "FFFFFF", True, "center")
    add_source(slide, "来源：https://arxiv.org/abs/2604.24890")
    return slide


def slide_industry(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_title(slide, "2026 工业落地：C2PA + SynthID 双层标识", "OpenAI 和 Google 的最新动作说明：metadata 与水印开始被组合使用。", "12 / INDUSTRY", 13)
    # Two layer diagram
    add_card(slide, 0.95, 1.65, 5.35, 3.9, "第一层：C2PA 内容凭证", [
        "提供模型、平台、时间、编辑历史",
        "信息量大，可验证签名",
        "问题：metadata 容易被剥离",
    ], "FFFFFF", "2563EB", 18, 13)
    add_card(slide, 7.0, 1.65, 5.35, 3.9, "第二层：SynthID 水印", [
        "嵌入不可见信号",
        "在截图/转换后可能更耐久",
        "问题：依赖生成方采用该体系",
    ], "FFFFFF", "0F766E", 18, 13)
    add_arrow(slide, 6.35, 3.60, 6.92, 3.60, "F59E0B", 2.2)
    add_badge(slide, "双层证据", 5.92, 3.08, 1.02, "F59E0B")
    add_text(slide, "对本项目的价值：现有“C2PA + 水印 + 来源归因”设计方向与 2026 产业路线一致。", 1.05, 6.28, 11.1, 0.28, 14, "111827", True, "center")
    add_source(slide, "来源：OpenAI / Google 2026 SynthID + C2PA updates")
    return slide


def slide_watermark_attacks(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, "F8FAFC")
    add_title(slide, "2026 水印安全：生产级水印也需要攻击测试", "文本水印在理论、改写攻击和供应链层面都出现新的安全讨论。", "13 / WATERMARK", 14)
    items = [
        ("SynthID-Text 分析", "Tournament sampling 的检测分数和层数设计会影响鲁棒性。", "2563EB"),
        ("Vaporizer", "翻译、同义替换、神经改写可削弱文本水印，同时保留语义。", "DC2626"),
        ("Blind PRNG Hijacking", "从随机源供应链层影响水印信号，提醒要关注生成链路完整性。", "F59E0B"),
    ]
    for i, (t, b, c) in enumerate(items):
        add_card(slide, 0.95 + i * 4.15, 1.82, 3.45, 3.55, t, b, "FFFFFF", c, 17, 12)
    add_text(slide, "汇报表达：水印是强证据，但不是单点答案；系统要把水印检测结果和被动检测、指纹、来源归因一起解释。", 1.0, 6.22, 11.3, 0.34, 14, "334155", True)
    add_source(slide, "来源：arXiv 2603.03410 / 2605.07481 / 2605.28632")
    return slide


def slide_fusion_arch(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_title(slide, "多证据融合架构：把 2026 调研落回系统", "最终报告不只输出概率，而是输出可追溯证据链和限制说明。", "14 / ARCHITECTURE", 15)
    nodes = [
        ("上传内容", "文本/图像/音频/视频", "64748B"),
        ("预处理与指纹", "metadata / SHA256 / pHash", "0F766E"),
        ("被动检测", "text / image / audio / video", "2563EB"),
        ("主动证据", "C2PA / SynthID / AudioSeal", "F59E0B"),
        ("来源归因", "Top-K model attribution", "7C3AED"),
        ("综合报告", "概率 + 来源 + 证据 + 限制", "DC2626"),
    ]
    for i, (t, b, c) in enumerate(nodes):
        x = 0.55 + i * 2.08
        add_card(slide, x, 2.35, 1.75, 1.20, t, b, "FFFFFF", c, 12, 8.7)
        if i < len(nodes) - 1:
            add_arrow(slide, x + 1.75, 2.95, x + 2.05, 2.95, "94A3B8", 1.3)
    add_card(slide, 1.05, 4.75, 11.2, 0.95, "可靠性分层", "强证据：C2PA 签名、官方水印命中  |  中证据：Top-K 来源归因、历史指纹匹配  |  弱证据：PPL、风格、视觉 artifact、频谱异常", "111827", "14B8A6", 15, 11)
    return slide


def slide_project_mapping(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, "F8FAFC")
    add_title(slide, "对现有项目的调整映射", "把调研结果转成可以继续开发的模块任务。", "15 / MAPPING", 16)
    quads = [
        ("Detection", ["真实 Hive/Sightengine provider", "多 provider ensemble", "跨模型/跨域评测"], "2563EB"),
        ("Provenance", ["c2patool wrapper", "AudioSeal / VideoSeal adapter", "C2PA 只作证据之一"], "0F766E"),
        ("Storage", ["保存原始 metadata", "保存 hash / pHash / Top-K", "记录 provider 版本"], "F59E0B"),
        ("Dashboard", ["展示证据链", "展示可疑帧/片段", "展示失败场景说明"], "DC2626"),
    ]
    for i, (t, items, c) in enumerate(quads):
        x = 0.85 + (i % 2) * 6.05
        y = 1.55 + (i // 2) * 2.25
        add_card(slide, x, y, 5.35, 1.70, t, items, "FFFFFF", c, 17, 12)
    add_text(slide, "下月开发优先级：先接真实 provider，再接 c2patool，随后扩展样本库并补水印 adapter。", 0.95, 6.38, 11.3, 0.28, 14, "111827", True, "center")
    return slide


def slide_roadmap(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_title(slide, "下阶段计划", "按“可演示、可验证、可扩展”的顺序推进。", "16 / NEXT", 17)
    weeks = [
        ("第 1 周", "真实检测 provider", "接入 Hive 或 Sightengine，替换 demo_api 主链路"),
        ("第 2 周", "官方 C2PA", "实现 c2patool wrapper，统一 provenance JSON 输出"),
        ("第 3 周", "样本库与归因", "生成多模型文本/图片样本，重训 Top-K 分类器"),
        ("第 4 周", "水印与展示", "AudioSeal/VideoSeal adapter，Dashboard 展示证据链"),
    ]
    x_start = 0.95
    for i, (wk, title, body) in enumerate(weeks):
        x = x_start + i * 3.0
        add_circle_label(slide, str(i + 1), x + 0.95, 1.65, 0.52, ["0F766E", "2563EB", "F59E0B", "DC2626"][i])
        add_card(slide, x, 2.45, 2.55, 2.45, f"{wk}\n{title}", body, "FFFFFF", ["0F766E", "2563EB", "F59E0B", "DC2626"][i], 15, 10.5)
        if i < 3:
            add_arrow(slide, x + 2.56, 3.68, x + 2.92, 3.68, "94A3B8", 1.2)
    add_badge(slide, "目标：形成一版可在服务器演示、能解释证据、可持续更新的 AIGC 检测与溯源系统", 1.55, 6.22, 10.2, "111827")
    return slide


def slide_references(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, "111827")
    add_title(slide, "参考来源", "本页保留 2026 重点来源，便于答辩时说明调研依据。", "17 / REFERENCES", 18, True)
    refs = [
        "Detecting AIGC on Social Media with Multi-modal Language Models, arXiv:2606.11200",
        "NTIRE 2026 Robust AI-Generated Image Detection in the Wild, arXiv:2604.11487",
        "Out-of-box AI-generated image detector benchmark, arXiv:2602.07814",
        "DetectRL-X multilingual real-world text detection, arXiv:2605.15518",
        "TSM-Bench Wikipedia editing detection, arXiv:2605.31113",
        "Detecting the Machine benchmark, arXiv:2603.17522",
        "C2PA specifications fall short, arXiv:2604.24890",
        "SynthID-Text analysis, arXiv:2603.03410",
        "Vaporizer watermark attack, arXiv:2605.07481",
        "Blind PRNG Hijacking, arXiv:2605.28632",
        "OpenAI / Google 2026 SynthID + C2PA industry updates",
    ]
    add_multiline(slide, refs, 0.9, 1.55, 11.7, 4.85, 11, "E5E7EB", bullet=False, gap=3)
    add_text(slide, "完整调研文档：sources/aigc_detection_provenance_2026_latest_addendum.md", 0.9, 6.65, 9.8, 0.22, 10, "94A3B8")
    return slide


def build():
    prs = Presentation()
    prs.slide_width = Inches(WIDE_W)
    prs.slide_height = Inches(WIDE_H)
    # Remove default first slide if present in some templates is not needed.
    slide_cover(prs)
    slide_overview(prs)
    slide_before_after(prs)
    slide_deployment(prs)
    slide_status(prs)
    slide_why_2026(prs)
    slide_image_2026(prs)
    slide_out_of_box(prs)
    slide_text_2026(prs)
    slide_social_vlm(prs)
    slide_samples(prs)
    slide_provenance_security(prs)
    slide_industry(prs)
    slide_watermark_attacks(prs)
    slide_fusion_arch(prs)
    slide_project_mapping(prs)
    slide_roadmap(prs)
    slide_references(prs)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
